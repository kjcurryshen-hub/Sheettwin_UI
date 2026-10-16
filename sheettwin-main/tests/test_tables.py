"""Table detection on layouts beyond 'one row of text labels over numbers'."""

import sqlite3

import pytest
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side

from sheetir import build_ir


@pytest.fixture
def run(tmp_path):
    cons = []

    def build(fill_sheet):
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        fill_sheet(ws)
        n = len(cons)
        wb.save(tmp_path / f"t{n}.xlsx")
        res = build_ir(tmp_path / f"t{n}.xlsx", tmp_path / f"t{n}.db")
        con = sqlite3.connect(res.db_path)
        con.row_factory = sqlite3.Row
        cons.append(con)
        return con
    yield build
    for con in cons:
        con.close()


def tables(con):
    return [dict(r) for r in con.execute("SELECT * FROM tables ORDER BY table_id")]


def fields(con, table_id=1):
    return [dict(r) for r in con.execute("SELECT * FROM fields WHERE table_id=? ORDER BY col", (table_id,))]


def put(ws, rows, start=1):
    for i, row in enumerate(rows):
        for j, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(start + i, j, v)


def test_year_headers(run):
    con = run(lambda ws: put(ws, [
        ["Region", 2023, 2024, 2025],
        ["North", 10, 12, 15],
        ["South", 8, 9, 11],
        ["East", 5, 7, 6],
    ]))
    (t,) = tables(con)
    assert (t["header_row"], t["row_start"], t["row_end"], t["key_col"]) == (1, 2, 4, 1)
    assert [f["header"] for f in fields(con)] == ["Region", "2023", "2024", "2025"]


def test_lone_number_next_to_label_is_not_a_header(run):
    # 'Discount | 2000' (an amount, not a year sequence) above more label/value pairs
    con = run(lambda ws: put(ws, [["Discount", 2000], ["Freight", 120], ["Tax", 0.05]]))
    assert tables(con) == []


def test_styled_all_text_table(run):
    def fill(ws):
        put(ws, [["Name", "Department", "Office"],
                 ["An", "Sales", "Taipei"],
                 ["Binh", "R&D", "Tainan"],
                 ["Chi", "Finance", "Taipei"]])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.border = Border(bottom=Side(style="thin"))
    (t,) = tables(run(fill))
    assert (t["header_row"], t["row_end"], t["key_col"]) == (1, 4, 1)


def test_unstyled_all_text_block_is_not_a_table(run):
    con = run(lambda ws: put(ws, [["Name", "Department"], ["An", "Sales"], ["Binh", "R&D"]]))
    assert tables(con) == []


def test_styled_numeric_header(run):
    def fill(ws):
        put(ws, [["Size", 10, 20, 50],
                 ["S", 1.0, 1.9, 4.5],
                 ["M", 1.2, 2.3, 5.4],
                 ["L", 1.5, 2.8, 6.6]])
        for c in ws[1]:
            c.fill = PatternFill("solid", fgColor="DDDDDD")
    con = run(fill)
    (t,) = tables(con)
    assert t["header_row"] == 1 and t["row_end"] == 4
    assert [f["name"] for f in fields(con)] == ["size", "_10", "_20", "_50"]


def test_two_row_header_with_merged_groups(run):
    def fill(ws):
        put(ws, [["Region", "2024", None, "2025", None],
                 [None, "Q1", "Q2", "Q1", "Q2"],
                 ["North", 1, 2, 3, 4],
                 ["South", 5, 6, 7, 8],
                 ["East", 9, 10, 11, 12]])
        ws.merge_cells("A1:A2")
        ws.merge_cells("B1:C1")
        ws.merge_cells("D1:E1")
    con = run(fill)
    (t,) = tables(con)
    assert (t["header_top"], t["header_row"], t["row_start"], t["row_end"], t["key_col"]) == (1, 2, 3, 5, 1)
    assert [f["header"] for f in fields(con)] == ["Region", "2024 / Q1", "2024 / Q2", "2025 / Q1", "2025 / Q2"]
    assert [f["name"] for f in fields(con)] == ["region", "_2024_q1", "_2024_q2", "_2025_q1", "_2025_q2"]
    sem = con.execute("SELECT semantic FROM cells WHERE cell_id='Data!C4'").fetchone()[0]
    assert sem == 'data["South"]._2024_q2'


def test_unmerged_group_labels_span_to_the_right(run):
    con = run(lambda ws: put(ws, [
        [None, "Actual", None, "Budget", None],
        ["Item", "Qty", "Amount", "Qty", "Amount"],
        ["Pen", 1, 10, 2, 20],
        ["Ink", 3, 30, 4, 40],
    ]))
    (t,) = tables(con)
    assert t["header_top"] == 1
    assert [f["header"] for f in fields(con)] == ["Item", "Actual / Qty", "Actual / Amount", "Budget / Qty", "Budget / Amount"]


def test_title_above_header_is_not_a_group_header(run):
    def fill(ws):
        put(ws, [["Sales report 2025"], [None], ["Item", "Qty", "Price"],
                 ["Pen", 1, 10], ["Ink", 3, 30]])
        ws.merge_cells("A1:C1")
    (t,) = tables(run(fill))
    assert (t["header_top"], t["header_row"]) == (None, 3)

    def fill_adjacent(ws):
        put(ws, [["Sales report 2025"], ["Item", "Qty", "Price"], ["Pen", 1, 10], ["Ink", 3, 30]])
        ws.merge_cells("A1:C1")
    (t,) = tables(run(fill_adjacent))
    assert (t["header_top"], t["header_row"]) == (None, 2)


def test_form_field_above_header_is_not_a_group_header(run):
    con = run(lambda ws: put(ws, [
        ["Customer", "ACME"],
        ["Item", "Qty", "Price"], ["Pen", 1, 10], ["Ink", 3, 30],
    ]))
    (t,) = tables(con)
    assert (t["header_top"], t["header_row"]) == (None, 2)
    assert [f["header"] for f in fields(con)] == ["Item", "Qty", "Price"]


def test_single_blank_row_inside_body(run):
    def fill(ws):
        put(ws, [["Item", "Qty", "Price", "Amount"],
                 ["Pen", 1, 10, None], ["Ink", 3, 30, None],
                 [None],
                 ["Pad", 2, 15, None], ["Cap", 5, 2, None]])
        for r in (2, 3, 5, 6):
            ws.cell(r, 4, f"=B{r}*C{r}")
    con = run(fill)
    (t,) = tables(con)
    assert (t["row_start"], t["row_end"]) == (2, 6)
    issues = con.execute("SELECT cell_id, kind FROM issues WHERE kind='missing_formula'").fetchall()
    assert issues == []  # the blank row is a gap, not a row missing its formula
    assert con.execute("SELECT semantic FROM cells WHERE cell_id='Data!D5'").fetchone()[0] == 'data["Pad"].amount'


def test_total_row_after_blank_row(run):
    def fill(ws):
        put(ws, [["Item", "Qty", "Price"], ["Pen", 1, 10], ["Ink", 3, 30], ["Pad", 2, 15],
                 [None], ["Total", "=SUM(B2:B4)", None]])
    (t,) = tables(run(fill))
    assert (t["row_end"], t["total_row"]) == (4, 6)


def test_blank_row_then_new_header_starts_a_new_table(run):
    con = run(lambda ws: put(ws, [
        ["Item", "Qty", "Price"], ["Pen", 1, 10], ["Ink", 3, 30],
        [None],
        ["Code", "Rate", "Limit"], ["A", 0.1, 100], ["B", 0.2, 200],
    ]))
    ts = tables(con)
    assert [(t["header_row"], t["row_end"]) for t in ts] == [(1, 3), (5, 7)]


def test_sections_separated_by_blank_rows_and_headings_stay_one_table(run):
    # statistical layout: a section heading in the stub column between blank rows
    con = run(lambda ws: put(ws, [
        [None, "2023", "2024"],
        [None, "$m", "$m"],
        ["Receipts", None, None],
        ["Taxes", 10, 12], ["Fees", 3, 4],
        [None],
        ["Payments", None, None],
        ["Wages", 7, 8], ["Rent", 2, 2],
    ]))
    (t,) = tables(con)
    assert (t["header_row"], t["col_start"], t["row_end"], t["key_col"]) == (1, 1, 9, None)
    assert [f["header"] for f in fields(con)][1:] == ["2023", "2024"]


def test_four_blank_rows_end_the_table(run):
    con = run(lambda ws: put(ws, [
        ["Item", "Qty", "Price"], ["Pen", 1, 10], ["Ink", 3, 30],
        [None], [None], [None], [None],
        ["Pad", 2, 15], ["Cap", 5, 2],
    ]))
    (t,) = tables(con)
    assert t["row_end"] == 3


def test_series_metadata_rows_belong_to_the_header(run):
    # time-series layout: descriptor rows above the series IDs, dates down the first column
    con = run(lambda ws: put(ws, [
        [None, "Wages; Total", "Wages; Public"],
        ["Unit", "Percent", "Percent"],
        ["Frequency", "Quarter", "Quarter"],
        ["Series ID", "A01", "A02"],
        ["Q1", 1.1, 0.9], ["Q2", 1.2, 1.0], ["Q3", 1.0, 0.8],
    ]))
    (t,) = tables(con)
    assert (t["header_top"], t["header_row"], t["row_end"]) == (1, 4, 7)


def test_stub_column_without_header_joins_the_table(run):
    con = run(lambda ws: put(ws, [
        [None, "Earners", "Income"],
        ["North", 10, 100], ["South", 8, 90], ["East", 5, 70],
    ]))
    (t,) = tables(con)
    assert (t["col_start"], t["key_col"]) == (1, 1)


def test_spacer_column_does_not_split_a_wide_table(run):
    con = run(lambda ws: put(ws, [
        ["Country", "Men", "Women", None, "Men", "Women"],
        ["FR", 1, 2, None, 3, 4], ["DE", 5, 6, None, 7, 8], ["IT", 9, 10, None, 11, 12],
    ]))
    (t,) = tables(con)
    assert (t["col_start"], t["col_end"]) == (1, 6)


def test_formula_columns_raise_confidence(run):
    def fill(formulas):
        def f(ws):
            put(ws, [["Item", "Grade", "Score"], ["Pen", "A", 1], ["Ink", "B", 0], ["Pad", "A", 1]])
            if formulas:
                for r in (2, 3, 4):
                    ws.cell(r, 3, f'=IF(B{r}="A",1,0)')
        return f
    (typed,) = tables(run(fill(False)))
    con = run(fill(True))
    (computed,) = tables(con)
    assert computed["confidence"] > typed["confidence"]
    assert fields(con)[2]["kind"] == "derived"


def test_confidence_is_stored(run):
    con = run(lambda ws: put(ws, [["Item", "Qty"], ["Pen", 1], ["Ink", 3]]))
    (t,) = tables(con)
    assert 0.7 <= t["confidence"] <= 1.0
