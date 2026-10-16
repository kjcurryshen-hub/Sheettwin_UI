import sqlite3

import pytest
from openpyxl import Workbook

from sheetir import build_ir
from sheetir.sample import PRODUCTS, make_quotation


def q(con, sql, *args):
    return con.execute(sql, args).fetchall()


@pytest.fixture(scope="module")
def con(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("rel")
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["Order", "Customer", "Qty", "Region", "Discount", "Tier", "Rate"])
    for i, (cust, qty) in enumerate([("ACME", 5), ("Globex", 12), ("Initech", 40)], start=2):
        ws.append([f"O-{i}", cust, qty,
                   f"=INDEX(Customers!$C$2:$C$4,MATCH(B{i},Customers!$A$2:$A$4,0))",
                   f"=VLOOKUP(C{i},Brackets!$A$2:$B$4,2)",          # approximate on numbers: legitimate
                   f"=VLOOKUP(B{i},Customers!$A$2:$C$4,2)",         # approximate on text: a bug
                   f"=XLOOKUP(B{i},Customers!$A$2:$A$4,Customers!$D$2:$D$4)"])

    cu = wb.create_sheet("Customers")
    cu.append(["Customer", "Tier", "Region", "Rate"])
    for row in [("ACME", "A", "North", 0.1), ("Globex", "B", "South", 0.2), ("Initech", "A", "North", 0.3)]:
        cu.append(row)
    cu["F1"], cu["G1"] = "Qty", "=SUMIFS(Orders!$C$2:$C$4,Orders!$B$2:$B$4,A2,Orders!$C$2:$C$4,\">10\")"
    cu["F2"], cu["G2"] = "Dynamic", '=VLOOKUP(A2,INDIRECT("Customers!A2:C4"),2,FALSE)'
    cu["F3"], cu["G3"] = "Header", "=HLOOKUP(\"Region\",$A$1:$D$4,2,FALSE)"

    br = wb.create_sheet("Brackets")
    br.append(["From", "Discount"])
    for row in [(0, 0), (10, 0.05), (30, 0.1)]:
        br.append(row)
    wb.save(tmp / "rel.xlsx")

    res = build_ir(tmp / "rel.xlsx", tmp / "rel.db")
    c = sqlite3.connect(res.db_path)
    yield c
    c.close()


def rel(con, cell, function):
    cols = "match, value, value_cell, key_range, return_range, from_field, key_field, return_field"
    rows = q(con, f"SELECT {cols} FROM relations WHERE cell_id = ? AND function = ?", cell, function)
    assert len(rows) == 1, rows
    return dict(zip(cols.split(", "), rows[0]))


def test_index_match_is_one_lookup(con):
    r = rel(con, "Orders!D2", "INDEX/MATCH")
    assert r["match"] == "exact"
    assert (r["key_range"], r["return_range"]) == ("Customers!A2:A4", "Customers!C2:C4")
    assert (r["from_field"], r["key_field"], r["return_field"]) == ("customer", "customer", "region")
    assert q(con, "SELECT COUNT(*) FROM relations WHERE cell_id='Orders!D2'") == [(1,)]  # MATCH not counted twice


def test_vlookup_match_modes(con):
    assert rel(con, "Orders!E2", "VLOOKUP")["match"] == "approximate"
    t = rel(con, "Orders!F2", "VLOOKUP")
    assert (t["match"], t["key_range"], t["return_range"]) == ("approximate", "Customers!A2:A4", "Customers!B2:B4")


def test_xlookup_defaults_to_exact(con):
    r = rel(con, "Orders!G2", "XLOOKUP")
    assert (r["match"], r["key_field"], r["return_field"]) == ("exact", "customer", "rate")


def test_sumifs_yields_one_relation_per_criterion(con):
    rows = q(con, "SELECT match, value, value_cell, key_range, return_range FROM relations "
                  "WHERE cell_id='Customers!G1' ORDER BY relation_id")
    assert rows == [("equal", "A2", "Customers!A2", "Orders!B2:B4", "Orders!C2:C4"),
                    ("condition", '">10"', None, "Orders!C2:C4", "Orders!C2:C4")]


def test_computed_range_is_recorded_without_a_key(con):
    r = rel(con, "Customers!G2", "VLOOKUP")
    assert (r["match"], r["key_range"]) == ("exact", None)


def test_hlookup_key_is_a_row(con):
    r = rel(con, "Customers!G3", "HLOOKUP")
    assert (r["key_range"], r["return_range"], r["key_field"]) == ("Customers!A1:D1", "Customers!A2:D2", None)


def test_approximate_lookup_flagged_only_on_text_keys(con):
    flagged = q(con, "SELECT cell_id, group_id FROM issues WHERE kind='approximate_lookup'")
    assert [c for c, _ in flagged] == ["Orders!F2"]
    assert flagged[0][1] is not None  # one issue for the whole filled-down group


def test_table_links_view_summarizes_joins(con):
    links = q(con, "SELECT from_table, from_field, to_table, key_field, return_field, function, n_cells "
                   "FROM table_links WHERE from_table='orders' ORDER BY function")
    assert ("orders", "customer", "customers", "customer", "region", "INDEX/MATCH", 3) in links
    assert ("orders", "customer", "customers", "customer", "rate", "XLOOKUP", 3) in links


def test_sample_material_lookup(tmp_path):
    res = build_ir(make_quotation(tmp_path / "s.xlsx"), tmp_path / "s.db")
    con = sqlite3.connect(res.db_path)
    assert q(con, "SELECT from_field, key_field, return_field, match, n_cells FROM table_links "
                  "WHERE from_table='quote' AND to_table='materials'") == [
        ("material", "code", "price_factor", "exact", len(PRODUCTS))]
    assert q(con, "SELECT COUNT(*) FROM issues WHERE kind='approximate_lookup'") == [(0,)]
