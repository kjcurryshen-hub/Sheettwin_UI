import sqlite3

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from sheetir import build_ir
from sheetir.sample import PRODUCTS, make_quotation
from sheetir.store import TRACE_SQL


@pytest.fixture(scope="module")
def ir(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("ir")
    res = build_ir(make_quotation(tmp / "quotation.xlsx"), tmp / "q.db")
    con = sqlite3.connect(res.db_path)
    yield res, con
    con.close()


def q(con, sql, *args):
    return con.execute(sql, args).fetchall()


def test_self_checks_pass(ir):
    res, _ = ir
    failed = [c for c in res.checks if not c[1] and c[2] == "error"]
    assert failed == []


def test_filled_down_formulas_collapse_into_groups(ir):
    _, con = ir
    (n_formulas,) = q(con, "SELECT COUNT(*) FROM cells WHERE vtype='formula'")[0]
    (n_groups,) = q(con, "SELECT COUNT(*) FROM formula_groups")[0]
    assert n_groups * 5 < n_formulas
    rows = q(con, "SELECT anchor, last, n_cells FROM formula_groups WHERE r1c1 = '=RC[-2]+RC[-1]'")
    assert rows == [("Quote!H2", f"Quote!H{1 + len(PRODUCTS)}", len(PRODUCTS))]


def test_quote_table_detected_with_key_total_and_fields(ir):
    _, con = ir
    (name, key_col, total_row), = q(con, "SELECT name, key_col, total_row FROM tables WHERE sheet='Quote'")
    assert (name, key_col, total_row) == ("quote", 1, 2 + len(PRODUCTS))
    fields = dict(q(con, "SELECT f.name, f.kind FROM fields f JOIN tables t USING(table_id) WHERE t.sheet='Quote'"))
    assert fields["copper"] == "data" and fields["margin"] == "derived"
    assert q(con, "SELECT unit FROM fields WHERE name='copper'") == [("kg",)]


def test_no_spurious_table_from_label_value_pairs(ir):
    _, con = ir
    assert q(con, "SELECT COUNT(*) FROM tables WHERE sheet IN ('Report','Inputs')") == [(0,)]


def test_params_named_from_labels_and_defined_names(ir):
    _, con = ir
    params = {n: (c, u) for c, n, u in q(con, "SELECT cell_id, name, unit FROM params")}
    assert params["copper_price"] == ("Inputs!B3", "TWD/kg")
    assert params["labor_rate"][0] == "Inputs!B5"   # from defined name LaborRate
    assert "unused_buffer" not in params            # orphan, not an input
    (domain,), = q(con, "SELECT domain FROM params WHERE name='copper_price'")
    assert '"between"' in domain


def test_semantic_addresses(ir):
    _, con = ir
    assert q(con, "SELECT semantic FROM cells WHERE cell_id='Quote!J2'") == [('quote["Motor housing A"].margin',)]
    assert q(con, "SELECT semantic FROM cells WHERE cell_id='Quote!I22'") == [("quote.total.price",)]
    assert q(con, "SELECT semantic FROM cells WHERE cell_id='Report!B3'") == [("outputs.lowest_margin",)]


def test_planted_problems_are_found(ir):
    _, con = ir
    found = {(c, k) for c, k in q(con, "SELECT cell_id, kind FROM issues")}
    assert ("Quote!I13", "inconsistent_formula") in found
    assert ("Quote!G17", "hardcoded_in_formula_column") in found
    assert ("Inputs!B9", "orphan") in found
    assert ("Report!B1", "volatile") in found
    assert ("Report!B6", "dynamic_ref") in found
    consts = q(con, "SELECT detail FROM issues WHERE kind='hardcoded_constant'")
    assert any("1.15" in d for (d,) in consts)
    assert not any("literal 3 " in d for (d,) in consts)  # VLOOKUP column index is not a business constant


def test_trace_reaches_inputs_through_name_and_cross_sheet_refs(ir):
    _, con = ir
    trace = {row[0] for row in q(con, TRACE_SQL, "Quote!J2", 10)}
    assert {"Quote!I2", "Quote!H2", "Inputs!B3", "Inputs!B5", "Inputs!B6", "Materials!C4"} <= trace
    assert "Quote!J3" not in trace


def test_listobject_and_structured_refs(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Item", "Qty", "Price", "Amount"])
    for i in range(2, 6):
        ws.append([f"item{i}", i, 10 * i, "=[@Qty]*[@Price]"])
    ws.add_table(Table(displayName="Sales", ref="A1:D5"))
    ws["F1"], ws["G1"] = "Grand total", "=SUM(Sales[Amount])"
    wb.save(tmp_path / "t.xlsx")

    res = build_ir(tmp_path / "t.xlsx", tmp_path / "t.db")
    con = sqlite3.connect(res.db_path)
    assert q(con, "SELECT name, source, key_col FROM tables") == [("Sales", "listobject", 1)]
    assert q(con, "SELECT n_cells FROM formula_groups WHERE anchor='Data!D2'") == [(4,)]
    trace = {row[0] for row in q(con, TRACE_SQL, "Data!G1", 5)}
    assert {"Data!D2", "Data!D5", "Data!B3", "Data!C4"} <= trace
    assert q(con, "SELECT COUNT(*) FROM issues WHERE kind LIKE 'unresolved%'") == [(0,)]


def test_cycle_is_reported(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = "=A2+1", "=A1+1", "=A2*2"
    wb.save(tmp_path / "c.xlsx")
    res = build_ir(tmp_path / "c.xlsx", tmp_path / "c.db")
    acyclic = next(c for c in res.checks if c[0] == "acyclic")
    assert not acyclic[1]
