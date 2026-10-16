import io
import sqlite3
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from sheetir.build import build_ir
from sheetir.fixes import _formula_block, _recalc_on_load, apply, suggest
from sheetir.sample import make_quotation


def _status(db):
    con = sqlite3.connect(db)
    try:
        return dict(con.execute("SELECT key, value FROM meta").fetchall())["parse_status"]
    finally:
        con.close()


def _fix_all(tmp_path, src):
    build_ir(src, tmp_path / "ir.db")
    s = suggest(tmp_path / "ir.db", src)
    edits = {f["cell"]: f["fix"] for f in s["fixes"] if f["fix"] and not f.get("blocked")}
    out = tmp_path / "fixed.xlsx"
    out.write_bytes(apply(src, edits))
    build_ir(out, tmp_path / "fixed.db")
    return s, out


def test_sample_indirect_becomes_a_direct_reference_and_verifies(tmp_path):
    src = make_quotation(tmp_path / "q.xlsx")
    s, out = _fix_all(tmp_path, src)
    assert _status(tmp_path / "ir.db") == "partial"
    (fix,) = s["fixes"]
    assert fix["cell"] == "Report!B6" and fix["fix"] == "=Quote!J5" and fix["inputs"] == [{"cell": "Report!B5", "value": "J5"}]
    assert _status(tmp_path / "fixed.db") == "verified"
    assert load_workbook(out)["Report"]["B6"].value == "=Quote!J5"


@pytest.fixture
def messy(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Calc"
    data = wb.create_sheet("Data")
    for r in range(1, 6):
        data.cell(r, 1, r * 10)
        data.cell(r, 2, r * 100)
    ws["A1"], ws["B1"] = "rows", 2
    ws["A2"] = "=OFFSET(Data!A1,B1,1)"            # -> Data!B3
    ws["A3"] = "=SUM(OFFSET(Data!A1,0,0,3,1))"    # -> SUM(Data!A1:A3)
    ws["A4"] = "=LaborRat*2"                      # typo of the name LaborRate
    ws["A5"] = "=Dta!A1+1"                        # typo of the sheet Data
    ws["A6"] = '=INDIRECT("R1C1",FALSE)'          # R1C1 style: not resolved
    ws["A7"] = "=INDIRECT(A8)"                    # A8 is a formula without a cached value
    ws["A8"] = '="Data!A1"'
    wb.defined_names["LaborRate"] = DefinedName("LaborRate", attr_text="Data!$A$5")
    p = tmp_path / "messy.xlsx"
    wb.save(p)
    return p


def test_offset_name_and_sheet_typos_are_fixed_the_rest_gets_advice(tmp_path, messy):
    s, out = _fix_all(tmp_path, messy)
    by = {f["cell"]: f for f in s["fixes"]}
    assert by["Calc!A2"]["fix"] == "=Data!B3" and by["Calc!A2"]["inputs"] == [{"cell": "Calc!B1", "value": 2}]
    assert by["Calc!A3"]["fix"] == "=SUM(Data!A1:A3)"
    assert by["Calc!A4"]["fix"] == "=LaborRate*2"
    assert by["Calc!A5"]["fix"] == "=Data!A1+1"
    assert by["Calc!A6"]["fix"] is None and by["Calc!A6"]["advice"] == ["dynamic_ref"]
    assert by["Calc!A7"]["fix"] is None
    ws = load_workbook(out)["Calc"]
    assert [ws[f"A{r}"].value for r in range(2, 6)] == ["=Data!B3", "=SUM(Data!A1:A3)", "=LaborRate*2", "=Data!A1+1"]
    con = sqlite3.connect(tmp_path / "fixed.db")
    left = {c for (c,) in con.execute("SELECT DISTINCT cell_id FROM issues WHERE kind IN "
                                      "('dynamic_ref','unresolved_name','unknown_sheet')")}
    con.close()
    assert left == {"Calc!A6", "Calc!A7"}


def test_apply_touches_only_the_edited_cells(tmp_path):
    src = make_quotation(tmp_path / "q.xlsx")
    with zipfile.ZipFile(src, "a") as z:  # a part openpyxl would drop on a re-save
        z.writestr("xl/media/image1.png", b"\x89PNG not really")
    out = apply(src, {"Report!B6": "=Quote!J5"})
    a, b = zipfile.ZipFile(src), zipfile.ZipFile(io.BytesIO(out))
    assert a.namelist() == b.namelist()
    changed = [n for n in a.namelist() if a.read(n) != b.read(n)]
    assert "xl/worksheets/sheet4.xml" in changed and set(changed) <= {"xl/workbook.xml", "xl/worksheets/sheet4.xml"}
    assert b'fullCalcOnLoad="1"' in b.read("xl/workbook.xml")


def test_shared_and_array_formulas():
    xml = ('<sheetData><row r="1"><c r="A1"><f t="shared" ref="A1:A3" si="0">B1*2</f><v>2</v></c></row>'
           '<row r="2"><c r="A2"><f t="shared" si="0"/><v>4</v></c></row>'
           '<row r="3"><c r="C3"><f t="array" ref="C3">SUM(B1:B2*2)</f></c></row></sheetData>')
    assert _formula_block(xml, "A1") == "shared_formula"
    assert _formula_block(xml, "C3") == "array_formula"
    assert _formula_block(xml, "Z9") == "not_found"
    c, f = _formula_block(xml, "A2")
    assert f.group(0) == '<f t="shared" si="0"/>'


def test_recalc_on_load_keeps_the_schema_order():
    wb = '<workbook><sheets><sheet name="A" sheetId="1" r:id="rId1"/></sheets><definedNames/><extLst/></workbook>'
    assert _recalc_on_load(wb) == ('<workbook><sheets><sheet name="A" sheetId="1" r:id="rId1"/></sheets><definedNames/>'
                                   '<calcPr fullCalcOnLoad="1"/><extLst/></workbook>')
    assert _recalc_on_load('<workbook><calcPr calcId="191029"/></workbook>') == \
        '<workbook><calcPr fullCalcOnLoad="1" calcId="191029"/></workbook>'


def test_apply_writes_values_new_cells_and_clears(tmp_path):
    src = make_quotation(tmp_path / "q.xlsx")
    out = tmp_path / "e.xlsx"
    out.write_bytes(apply(src, {"Inputs!C2": "between", "Inputs!B40": "7.5", "Inputs!A40": "R&D <factor>",
                                "Inputs!B9": "", "Inputs!D3": "TRUE", "Quote!G17": "=E17*LaborRate"}))
    wb = load_workbook(out)
    ws = wb["Inputs"]
    assert (ws["C2"].value, ws["B40"].value, ws["A40"].value, ws["B9"].value, ws["D3"].value) == \
        ("between", 7.5, "R&D <factor>", None, True)
    assert wb["Quote"]["G17"].value == "=E17*LaborRate"
    assert ws.max_row == 40
    build_ir(out, tmp_path / "e.db")
    assert _status(tmp_path / "e.db") == "partial"  # Report!B6 still uses INDIRECT; the check itself passes
