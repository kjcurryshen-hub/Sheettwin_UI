"""Turn a .csv, .xls or .ods file into an .xlsx, so the rest of SheetIR (and its independent check) reads one format.

The format is recognized from the file's bytes, not its name. What each conversion keeps:

  .csv  one sheet of values; numbers, booleans, ISO dates and percentages are typed, everything else is text
        (warning "from_csv")
  .xls  values, cached results of formulas, merged cells, bold / fill / bottom border, number formats, defined
        names; the formulas themselves are not readable by xlrd and are lost (warnings "from_xls", "formulas_dropped")
  .ods  values, formulas (translated from OpenFormula) with their cached results, merged cells, bold / fill / bottom
        border, named ranges (warning "from_ods")
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.workbook.defined_name import DefinedName

from .loader import WorkbookError

SUFFIXES = (".xlsx", ".xlsm", ".csv", ".tsv", ".xls", ".ods")

_OLE2 = bytes.fromhex("d0cf11e0a1b11ae1")
_MAX_ROW, _MAX_COL = 1_048_576, 16_384
_ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}
_BAD_TITLE = re.compile(r"[\[\]:*?/\\]")


def as_xlsx(src: Path, workdir: Path) -> tuple[Path, list[str]]:
    """The workbook to parse and the warnings its conversion raised: `src` itself when it is already an .xlsx /
    .xlsm (or not recognized: the loader then explains why), otherwise an .xlsx written into `workdir`."""
    src = Path(src)
    data = src.read_bytes()
    if data.startswith(_OLE2):
        encrypted = "EncryptionInfo".encode("utf-16-le") in data or "EncryptedPackage".encode("utf-16-le") in data
        if encrypted:
            return src, []  # an encrypted .xlsx: the loader reports it
        book, warnings = _from_xls(data), ["from_xls", "formulas_dropped"]
    elif zipfile.is_zipfile(io.BytesIO(data)):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = set(z.namelist())
            mimetype = z.read("mimetype").strip() if "mimetype" in names else b""
        if mimetype != b"application/vnd.oasis.opendocument.spreadsheet":
            return src, []
        book, warnings = _from_ods(data), ["from_ods"]
    elif _looks_like_text(data):
        book, warnings = _from_csv(data, src), ["from_csv"]
    else:
        return src, []
    out = Path(workdir) / (src.stem + ".xlsx")
    book.save(out)
    return out, warnings


# ---- writing ----------------------------------------------------------------------------------------------------


class _Book:
    """Sheets of cells to write: per cell a constant or a formula with its cached result, and header-like marks."""

    def __init__(self):
        self.sheets: list[dict] = []
        self.names: list[tuple[str, str]] = []

    def sheet(self, title: str) -> dict:
        title = _BAD_TITLE.sub("_", title).strip("'")[:31] or "Sheet"
        base, n = title, 2
        while any(s["title"].lower() == title.lower() for s in self.sheets):
            title = f"{base[:28]}({n})"
            n += 1
        s = {"title": title, "cells": {}, "merged": []}
        self.sheets.append(s)
        return s

    def save(self, path: Path) -> None:
        wb = Workbook()
        wb.remove(wb.active)
        cached: dict[int, dict[tuple[int, int], object]] = {}
        for i, s in enumerate(self.sheets):
            ws = wb.create_sheet(s["title"])
            for (r, c), (value, formula, marks, fmt) in s["cells"].items():
                cell = ws.cell(r, c)
                if formula is not None:
                    cell.value = formula
                    cached.setdefault(i, {})[(r, c)] = value
                else:
                    cell.value = value
                    if isinstance(value, str) and value.startswith("="):
                        cell.data_type = "s"  # text that only looks like a formula
                if fmt:
                    cell.number_format = fmt
                if "bold" in marks:
                    cell.font = Font(bold=True)
                if "fill" in marks:
                    cell.fill = PatternFill("solid", fgColor="DDDDDD")
                if "border_bottom" in marks:
                    cell.border = Border(bottom=Side(style="thin"))
            for r1, c1, r2, c2 in s["merged"]:
                ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        for name, ref in self.names:
            try:
                wb.defined_names[name] = DefinedName(name, attr_text=ref)
            except Exception:  # noqa: BLE001 - a name Excel would not accept: leave it out
                pass
        if not wb.worksheets:
            wb.create_sheet("Sheet")
        wb.save(path)
        if cached:
            _write_cached(path, cached)


_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _write_cached(path: Path, cached: dict[int, dict[tuple[int, int], object]]) -> None:
    """openpyxl writes formulas without results; put the results the source file had next to them."""
    ET.register_namespace("", _NS)
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    with zipfile.ZipFile(path) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    for i, values in cached.items():
        name = f"xl/worksheets/sheet{i + 1}.xml"
        root = ET.fromstring(parts[name])
        for c in root.iter(f"{{{_NS}}}c"):
            f = c.find(f"{{{_NS}}}f")
            if f is None:
                continue
            col = re.match(r"[A-Z]+", c.get("r")).group()
            row = int(c.get("r")[len(col):])
            value = values.get((row, _col_index(col)))
            for v in c.findall(f"{{{_NS}}}v"):
                c.remove(v)
            if value is None:
                continue
            c.attrib.pop("t", None)
            if isinstance(value, bool):
                c.set("t", "b")
                text = "1" if value else "0"
            elif isinstance(value, (int, float)):
                text = repr(value)
            elif isinstance(value, (dt.datetime, dt.date, dt.time)):
                text = repr(to_excel(value))
            elif isinstance(value, dt.timedelta):
                text = repr(value.total_seconds() / 86400)
            else:
                c.set("t", "e" if value in _ERRORS else "str")
                text = str(value)
            ET.SubElement(c, f"{{{_NS}}}v").text = text
        parts[name] = ET.tostring(root, xml_declaration=True, encoding="UTF-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


# ---- .csv -------------------------------------------------------------------------------------------------------


def _looks_like_text(data: bytes) -> bool:
    if not data.strip():
        return False
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return True
    head = data[:4096].lstrip().lower()
    return b"\0" not in data[:4096] and not head.startswith((b"<html", b"<!doctype", b"<?xml", b"%pdf"))


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for enc in ("utf-8-sig", "cp950", "gb18030"):  # cp950: Big5 as saved by Excel in Taiwan
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


_INT = re.compile(r"[+-]?(0|[1-9]\d{0,14})")
_FLOAT = re.compile(r"[+-]?(\d+\.\d*|\.\d+|\d+)([eE][+-]?\d+)?")
_GROUPED = re.compile(r"[+-]?\d{1,3}(,\d{3})+(\.\d+)?")
_PERCENT = re.compile(r"([+-]?(\d+\.?\d*|\.\d+))%")
_DATE = re.compile(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?)?")


def _typed(text: str, delimiter: str) -> tuple[object, str | None]:
    """The value Excel would see in a CSV field (and a number format for it). Leading zeros stay text."""
    s = text.strip()
    if _INT.fullmatch(s):
        return int(s), None
    if _FLOAT.fullmatch(s) and not re.fullmatch(r"[+-]?(0\d+|\d{16,})", s.split(".")[0].split("e")[0].split("E")[0]):
        # leading zeros or more digits than Excel keeps: ids, phone numbers, account numbers stay text
        return float(s), None
    if delimiter != "," and _GROUPED.fullmatch(s):
        v = float(s.replace(",", ""))
        return (int(v) if v.is_integer() and "." not in s else v), "#,##0" if "." not in s else "#,##0.00"
    if m := _PERCENT.fullmatch(s):
        return float(m.group(1)) / 100, "0%" if "." not in m.group(1) else "0.0%"
    if s.upper() in ("TRUE", "FALSE"):
        return s.upper() == "TRUE", None
    if m := _DATE.fullmatch(s):
        y, mo, d, h, mi, sec = (int(g) if g else 0 for g in m.groups())
        try:
            value = dt.datetime(y, mo, d, h, mi, sec)
        except ValueError:
            return text, None
        return value, "yyyy-mm-dd hh:mm:ss" if m.group(4) else "yyyy-mm-dd"
    return text, None


def _from_csv(data: bytes, src: Path) -> _Book:
    text = _decode(data)
    sample = text[:65536]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = "\t" if src.suffix.lower() == ".tsv" else ","
    book = _Book()
    s = book.sheet(src.stem)
    for r, row in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter), start=1):
        if r > _MAX_ROW:
            break
        for c, field in enumerate(row[:_MAX_COL], start=1):
            if field.strip() == "":
                continue
            value, fmt = _typed(field, delimiter)
            s["cells"][(r, c)] = (value, None, (), fmt)
    return book


# ---- .xls -------------------------------------------------------------------------------------------------------


def _from_xls(data: bytes) -> _Book:
    import xlrd

    try:
        try:
            wb = xlrd.open_workbook(file_contents=data, formatting_info=True)
        except NotImplementedError:  # formatting of some files is not supported: read the values alone
            wb = xlrd.open_workbook(file_contents=data)
    except xlrd.XLRDError as exc:
        raise WorkbookError("encrypted" if "encrypt" in str(exc).lower() else "corrupt", str(exc)) from exc
    except Exception as exc:  # noqa: BLE001 - anything xlrd raises means the file cannot be read
        raise WorkbookError("corrupt", f"{type(exc).__name__}: {exc}") from exc

    book = _Book()
    titles = []
    for sh in wb.sheets():
        s = book.sheet(sh.name)
        titles.append(s["title"])
        for r in range(min(sh.nrows, _MAX_ROW)):
            for c in range(min(sh.ncols, _MAX_COL)):
                cell = sh.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                if cell.ctype == xlrd.XL_CELL_NUMBER:
                    value = int(cell.value) if float(cell.value).is_integer() else cell.value
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(cell.value, wb.datemode)
                    except Exception:  # noqa: BLE001 - an out-of-range date: keep the serial number
                        value = cell.value
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(cell.value)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    value = xlrd.error_text_from_code.get(cell.value, "#VALUE!")
                else:
                    value = cell.value
                    if value == "":
                        continue
                marks, fmt = _xls_style(wb, sh, r, c)
                s["cells"][(r + 1, c + 1)] = (value, None, marks, fmt)
        for r1, r2, c1, c2 in getattr(sh, "merged_cells", []):
            if r2 - r1 > 1 or c2 - c1 > 1:
                s["merged"].append((r1 + 1, c1 + 1, r2, c2))
    for n in wb.name_obj_list:
        if n.builtin or n.hidden:
            continue
        try:
            ref = n.result.value[0]
            sheet = titles[ref.shtxlo]
            a = f"${get_column_letter(ref.colxlo + 1)}${ref.rowxlo + 1}"
            b = f"${get_column_letter(ref.colxhi)}${ref.rowxhi}"
            book.names.append((n.name, f"{_quote(sheet)}!{a}" + (f":{b}" if a != b else "")))
        except Exception:  # noqa: BLE001 - a name that is not a plain cell range
            continue
    return book


def _xls_style(wb, sh, r: int, c: int) -> tuple[frozenset[str], str | None]:
    if not wb.xf_list:
        return frozenset(), None
    xf = wb.xf_list[sh.cell_xf_index(r, c)]
    marks = set()
    if wb.font_list[xf.font_index].bold:
        marks.add("bold")
    if xf.background.fill_pattern:
        marks.add("fill")
    if xf.border.bottom_line_style:
        marks.add("border_bottom")
    fmt = wb.format_map.get(xf.format_key)
    fmt = fmt.format_str if fmt is not None else None
    return frozenset(marks), (fmt if fmt and fmt != "General" else None)


# ---- .ods -------------------------------------------------------------------------------------------------------

_O = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
_T = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
_X = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
_S = "{urn:oasis:names:tc:opendocument:xmlns:style:1.0}"
_F = "{urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0}"


def _from_ods(data: bytes) -> _Book:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            content = ET.fromstring(z.read("content.xml"))
            styles = ET.fromstring(z.read("styles.xml")) if "styles.xml" in z.namelist() else None
    except Exception as exc:  # noqa: BLE001
        raise WorkbookError("corrupt", f"{type(exc).__name__}: {exc}") from exc

    marks_of = _ods_styles([x for x in (styles, content) if x is not None])
    book = _Book()
    for table in content.iter(f"{_T}table"):
        s = book.sheet(table.get(f"{_T}name") or "Sheet")
        col_styles = []
        for col in table.iter(f"{_T}table-column"):
            n = int(col.get(f"{_T}number-columns-repeated", "1"))
            col_styles.extend([col.get(f"{_T}default-cell-style-name")] * min(n, _MAX_COL - len(col_styles)))
        r = 0
        for row in table.iter(f"{_T}table-row"):
            rows = int(row.get(f"{_T}number-rows-repeated", "1"))
            cells = [ch for ch in row if ch.tag in (f"{_T}table-cell", f"{_T}covered-table-cell")]
            if not any(_ods_has_content(ch) for ch in cells):
                r += rows
                continue
            for _ in range(rows):
                r += 1
                if r > _MAX_ROW:
                    break
                c = 0
                for ch in cells:
                    cols = int(ch.get(f"{_T}number-columns-repeated", "1"))
                    if ch.tag == f"{_T}covered-table-cell" or not _ods_has_content(ch):
                        c += cols
                        continue
                    for _ in range(cols):
                        c += 1
                        if c > _MAX_COL:
                            break
                        style = ch.get(f"{_T}style-name") or (col_styles[c - 1] if c <= len(col_styles) else None)
                        value, formula, fmt = _ods_cell(ch)
                        if value is not None or formula is not None:
                            s["cells"][(r, c)] = (value, formula, marks_of.get(style, frozenset()), fmt)
                        span_c = int(ch.get(f"{_T}number-columns-spanned", "1"))
                        span_r = int(ch.get(f"{_T}number-rows-spanned", "1"))
                        if span_c > 1 or span_r > 1:
                            s["merged"].append((r, c, r + span_r - 1, c + span_c - 1))
    for nr in content.iter(f"{_T}named-range"):
        name, address = nr.get(f"{_T}name"), nr.get(f"{_T}cell-range-address")
        if name and address:
            book.names.append((name, _odf_ref(address, absolute=True)))
    return book


def _ods_has_content(cell) -> bool:
    return (cell.get(f"{_O}value-type") is not None or cell.get(f"{_T}formula") is not None
            or cell.find(f"{_X}p") is not None
            or int(cell.get(f"{_T}number-columns-spanned", "1")) > 1
            or int(cell.get(f"{_T}number-rows-spanned", "1")) > 1)


def _ods_text(el) -> str:
    out = [el.text or ""]
    for ch in el:
        if ch.tag == f"{_X}s":
            out.append(" " * int(ch.get(f"{_X}c", "1")))
        elif ch.tag == f"{_X}tab":
            out.append("\t")
        elif ch.tag == f"{_X}line-break":
            out.append("\n")
        else:
            out.append(_ods_text(ch))
        out.append(ch.tail or "")
    return "".join(out)


_DURATION = re.compile(r"-?PT?(?:(\d+)H)?(?:(\d+)M)?(?:([\d.]+)S)?")


def _ods_cell(cell) -> tuple[object, str | None, str | None]:
    """(value or cached result, formula in Excel syntax, number format)."""
    kind = cell.get(f"{_O}value-type")
    text = "\n".join(_ods_text(p) for p in cell.findall(f"{_X}p"))
    fmt = None
    if kind in ("float", "percentage", "currency"):
        value = float(cell.get(f"{_O}value", "0"))
        value = int(value) if value.is_integer() else value
        fmt = "0%" if kind == "percentage" else None
    elif kind == "date":
        raw = cell.get(f"{_O}date-value", "")
        try:
            value = dt.datetime.fromisoformat(raw)
        except ValueError:
            value = text
        fmt = "yyyy-mm-dd" if "T" not in raw else "yyyy-mm-dd hh:mm:ss"
    elif kind == "time":
        m = _DURATION.fullmatch(cell.get(f"{_O}time-value", ""))
        if m:
            h, mi, sec = int(m.group(1) or 0), int(m.group(2) or 0), float(m.group(3) or 0)
            value = dt.timedelta(hours=h, minutes=mi, seconds=sec)
            if value < dt.timedelta(days=1):
                value = (dt.datetime.min + value).time()
            fmt = "hh:mm:ss"
        else:
            value = text
    elif kind == "boolean":
        value = cell.get(f"{_O}boolean-value") == "true"
    else:
        value = text if text != "" else None
    formula = cell.get(f"{_T}formula")
    return value, (_odf_formula(formula) if formula else None), fmt


def _ods_styles(roots) -> dict[str, frozenset[str]]:
    """Cell style name -> header-like marks, following parent styles."""
    own: dict[str, tuple[set[str], str | None]] = {}
    for root in roots:
        for st in root.iter(f"{_S}style"):
            if st.get(f"{_S}family") != "table-cell":
                continue
            marks = set()
            tp = st.find(f"{_S}text-properties")
            if tp is not None and tp.get(f"{_F}font-weight") in ("bold", "600", "700", "800", "900"):
                marks.add("bold")
            cp = st.find(f"{_S}table-cell-properties")
            if cp is not None:
                bg = cp.get(f"{_F}background-color")
                if bg and bg != "transparent" and bg.lower() != "#ffffff":
                    marks.add("fill")
                border = cp.get(f"{_F}border-bottom") or cp.get(f"{_F}border")
                if border and border != "none" and not border.startswith("0"):
                    marks.add("border_bottom")
            own[st.get(f"{_S}name")] = (marks, st.get(f"{_S}parent-style-name"))
    out = {}
    for name in own:
        marks, seen, cur = set(), set(), name
        while cur in own and cur not in seen:
            seen.add(cur)
            marks |= own[cur][0]
            cur = own[cur][1]
        out[name] = frozenset(marks)
    return out


_ODF_REF = re.compile(r"\[([^\]]*)\]")


def _odf_formula(f: str) -> str:
    """OpenFormula as LibreOffice writes it ("of:=SUM([.A1:.B2];[Sheet2.C3])") in Excel syntax."""
    prefix = re.match(r"[a-z]+:", f)
    if prefix:
        if prefix.group() == "msoxl:":
            return f[len("msoxl:"):]
        f = f[prefix.end():]
    parts = re.split(r'("(?:[^"]|"")*")', f)  # leave string literals alone
    for i in range(0, len(parts), 2):
        parts[i] = _ODF_REF.sub(lambda m: _odf_ref(m.group(1)), parts[i]).replace(";", ",")
    return "".join(parts)


def _odf_ref(s: str, absolute: bool = False) -> str:
    """"$'My Sheet'.$A$1:.$B$2" -> "'My Sheet'!$A$1:$B$2" (".A1" -> "A1")."""
    if s.startswith("#"):
        return "#REF!"
    out, first = [], None
    for k, part in enumerate(s.split(":")):
        sheet, _, cell = part.rpartition(".")
        sheet = sheet.lstrip("$")
        if sheet.startswith("'") and sheet.endswith("'"):
            sheet = sheet[1:-1].replace("''", "'")
        if k == 0:
            first = sheet or None
        elif sheet and sheet != first:
            cell = f"{_quote(sheet)}!{cell}"
        out.append(cell)
    ref = ":".join(out)
    return f"{_quote(first)}!{ref}" if first else ref


def _quote(sheet: str) -> str:
    return sheet if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", sheet) else "'" + sheet.replace("'", "''") + "'"
