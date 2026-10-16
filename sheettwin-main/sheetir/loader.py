"""L0: read the raw workbook (formulas and cached values) without modifying it."""

from __future__ import annotations

import datetime as dt
import hashlib
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

CellKey = tuple[str, int, int]  # (sheet, row, col)


@dataclass
class RawCell:
    sheet: str
    row: int
    col: int
    vtype: str  # formula | number | text | bool | date | error
    value: object  # constant value; None for formulas
    formula: str | None
    cached: object  # value Excel last computed (None if never recalculated)
    number_format: str
    flags: list[str] = field(default_factory=list)
    style: frozenset[str] = frozenset()  # header-like formatting: bold | fill | border_bottom


@dataclass
class TableDef:
    """An Excel Table (ListObject) declared in the file."""

    name: str
    sheet: str
    ref: str
    header_rows: int
    totals_rows: int
    columns: list[str]


@dataclass
class DefinedName:
    name: str
    scope: str | None  # sheet name for sheet-scoped names
    refers_to: str


@dataclass
class Validation:
    sheet: str
    sqref: str
    rule: dict


@dataclass
class RawWorkbook:
    path: Path
    sha256: str
    sheets: list[str]
    cells: dict[CellKey, RawCell]
    names: list[DefinedName]
    tables: list[TableDef]
    validations: list[Validation]
    merged: list[tuple[str, str]]
    warnings: list[str] = field(default_factory=list)  # e.g. "styles_dropped": formatting unreadable, ignored


class WorkbookError(Exception):
    """The file cannot be read as a workbook. `kind` says why, so a caller can explain it and suggest a next step:
    legacy_xls (Excel 97-2003 format), encrypted (password protected), not_a_workbook, corrupt."""

    def __init__(self, kind: str, detail: str = ""):
        super().__init__(f"{kind}: {detail}" if detail else kind)
        self.kind, self.detail = kind, detail


_OLE2 = bytes.fromhex("d0cf11e0a1b11ae1")  # compound file: .xls, or an .xlsx encrypted with a password
_MINIMAL_STYLES = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    b'<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    b'<fonts count="1"><font/></fonts><fills count="2"><fill><patternFill patternType="none"/></fill>'
    b'<fill><patternFill patternType="gray125"/></fill></fills><borders count="1"><border/></borders>'
    b'<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
    b'<cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs></styleSheet>'
)


def _open(path: Path):
    """Both openpyxl views of the workbook (formulas, cached values) and any warnings.

    When the file opens but its formatting part is unreadable, the workbook is read again with that part replaced by
    a minimal one: cells and formulas are kept, formatting (bold, fills, borders) is lost (warning "styles_dropped")."""
    data = path.read_bytes()
    if data.startswith(_OLE2):
        encrypted = "EncryptionInfo".encode("utf-16-le") in data or "EncryptedPackage".encode("utf-16-le") in data
        raise WorkbookError("encrypted" if encrypted else "legacy_xls")
    if not zipfile.is_zipfile(path):
        raise WorkbookError("not_a_workbook")
    try:
        return load_workbook(path, data_only=False), load_workbook(path, data_only=True), []
    except Exception as exc:  # noqa: BLE001 - anything openpyxl raises means the same thing to the user
        first = exc
    try:
        with tempfile.TemporaryDirectory() as tmp:
            fixed = Path(tmp) / "styles_dropped.xlsx"
            with zipfile.ZipFile(path) as zin, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    zout.writestr(item, _MINIMAL_STYLES if item.filename == "xl/styles.xml" else zin.read(item.filename))
            return load_workbook(fixed, data_only=False), load_workbook(fixed, data_only=True), ["styles_dropped"]
    except Exception:  # noqa: BLE001
        raise WorkbookError("corrupt", f"{type(first).__name__}: {first}") from first


def _vtype(cell) -> str:
    if cell.data_type == "f":
        return "formula"
    if cell.data_type == "e":
        return "error"
    if cell.data_type == "b" or isinstance(cell.value, bool):
        return "bool"
    if cell.data_type == "d" or isinstance(cell.value, (dt.date, dt.datetime, dt.time)):
        return "date"
    if cell.data_type == "n":
        return "number"
    return "text"


def _style(cell) -> frozenset[str]:
    if not cell.has_style:
        return frozenset()
    marks = set()
    if cell.font is not None and cell.font.b:
        marks.add("bold")
    if cell.fill is not None and cell.fill.fill_type not in (None, "none"):
        marks.add("fill")
    if cell.border is not None and cell.border.bottom is not None and cell.border.bottom.style:
        marks.add("border_bottom")
    return frozenset(marks)


def load(path: str | Path) -> RawWorkbook:
    path = Path(path)
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    wb_f, wb_v, warnings = _open(path)

    cells: dict[CellKey, RawCell] = {}
    tables: list[TableDef] = []
    validations: list[Validation] = []
    merged: list[tuple[str, str]] = []
    names: list[DefinedName] = [DefinedName(n, None, d.attr_text) for n, d in wb_f.defined_names.items()]

    for ws in wb_f.worksheets:
        ws_v = wb_v[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                flags: list[str] = []
                value = cell.value
                formula = None
                if isinstance(value, ArrayFormula):
                    formula, flags = value.text, ["array_formula"]
                elif isinstance(value, DataTableFormula):
                    formula, flags = "=" + str(value), ["data_table"]
                elif cell.data_type == "f":
                    formula = value
                vtype = "formula" if formula is not None else _vtype(cell)
                cells[(ws.title, cell.row, cell.column)] = RawCell(
                    sheet=ws.title,
                    row=cell.row,
                    col=cell.column,
                    vtype=vtype,
                    value=None if formula is not None else value,
                    formula=formula,
                    cached=ws_v.cell(cell.row, cell.column).value if formula is not None else value,
                    number_format=cell.number_format,
                    flags=flags,
                    style=_style(cell),
                )

        for t in ws.tables.values():
            tables.append(TableDef(
                name=t.displayName or t.name,
                sheet=ws.title,
                ref=t.ref,
                header_rows=1 if t.headerRowCount is None else t.headerRowCount,
                totals_rows=t.totalsRowCount or 0,
                columns=[c.name for c in t.tableColumns],
            ))
        for dv in ws.data_validations.dataValidation:
            rule = {"type": dv.type, "operator": dv.operator, "formula1": dv.formula1, "formula2": dv.formula2}
            validations.append(Validation(ws.title, str(dv.sqref), rule))
        merged.extend((ws.title, str(r)) for r in ws.merged_cells.ranges)
        names.extend(DefinedName(n, ws.title, d.attr_text) for n, d in ws.defined_names.items())

    return RawWorkbook(path, sha, wb_f.sheetnames, cells, names, tables, validations, merged, warnings)
