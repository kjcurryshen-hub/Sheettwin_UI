"""Synthetic quotation workbook with deliberately planted problems, for demos and tests.

Planted:
  - Quote!I13   overhead 1.12 instead of 1.15        -> inconsistent_formula
  - Quote!G17   labor cost typed in as 3250           -> hardcoded_in_formula_column
  - 1.15        hard-coded overhead factor            -> hardcoded_constant
  - Inputs!B9   value nobody uses                     -> orphan
  - Report!B1   =TODAY()                              -> volatile
  - Report!B6   =INDIRECT(...)                        -> dynamic_ref
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

PRODUCTS = [
    # name, material code, copper kg, steel kg, labor hours
    ("Motor housing A", "CU", 1.8, 4.0, 2.5),
    ("Motor housing B", "CU", 2.4, 5.5, 3.0),
    ("Bracket S", "ST", 0.0, 1.2, 0.6),
    ("Bracket L", "ST", 0.0, 2.8, 0.9),
    ("Heat sink 40", "AL", 0.3, 0.0, 0.8),
    ("Heat sink 60", "AL", 0.5, 0.0, 1.1),
    ("Busbar 100A", "CU", 1.1, 0.0, 0.7),
    ("Busbar 250A", "CU", 2.9, 0.0, 1.0),
    ("Terminal block", "CU", 0.4, 0.2, 0.5),
    ("Cable lug kit", "CU", 0.6, 0.0, 0.4),
    ("Frame panel", "ST", 0.0, 6.5, 1.8),
    ("Cover plate", "ST", 0.0, 1.9, 0.5),
    ("Coil assembly", "CU", 3.5, 0.8, 4.2),
    ("Mounting rail", "ST", 0.0, 3.1, 0.6),
    ("Shield can", "AL", 0.2, 0.0, 0.7),
    ("Contact spring", "CU", 0.15, 0.05, 0.3),
    ("Gear housing", "AL", 0.4, 1.0, 2.2),
    ("Base plate", "ST", 0.0, 8.0, 1.5),
    ("Ground strap", "CU", 0.8, 0.0, 0.3),
    ("Fan bracket", "ST", 0.0, 0.9, 0.4),
]


def make_quotation(path: str | Path) -> Path:
    path = Path(path)
    wb = Workbook()

    inp = wb.active
    inp.title = "Inputs"
    inp["A1"] = "Quotation inputs"
    rows = [
        (3, "Copper price (TWD/kg)", 285),
        (4, "Steel price (TWD/kg)", 32),
        (5, "Labor rate (TWD/h)", 650),
        (6, "Freight per unit (TWD)", 120),
        (7, "Target margin", 0.15),
        (9, "Unused buffer", 0.05),
    ]
    for r, label, value in rows:
        inp.cell(r, 1, label)
        inp.cell(r, 2, value)
    inp["B7"].number_format = "0%"
    dv = DataValidation(type="decimal", operator="between", formula1="100", formula2="1000")
    dv.add("B3")
    inp.add_data_validation(dv)
    wb.defined_names["LaborRate"] = DefinedName("LaborRate", attr_text="Inputs!$B$5")

    mat = wb.create_sheet("Materials")
    for r, row in enumerate([("Code", "Description", "Price factor"),
                             ("CU", "Copper alloy", 1.00),
                             ("ST", "Steel", 1.00),
                             ("AL", "Aluminium", 1.35)], start=1):
        for c, v in enumerate(row, start=1):
            mat.cell(r, c, v)

    q = wb.create_sheet("Quote")
    headers = ["Product", "Material", "Copper (kg)", "Steel (kg)", "Labor (h)",
               "Material cost", "Labor cost", "Unit cost", "Price", "Margin"]
    for c, h in enumerate(headers, start=1):
        q.cell(1, c, h)
    first, last = 2, 1 + len(PRODUCTS)
    for i, (name, code, cu, st, lab) in enumerate(PRODUCTS):
        r = first + i
        q.cell(r, 1, name)
        q.cell(r, 2, code)
        q.cell(r, 3, cu)
        q.cell(r, 4, st)
        q.cell(r, 5, lab)
        q.cell(r, 6, f"=C{r}*Inputs!$B$3*VLOOKUP(B{r},Materials!$A$2:$C$4,3,FALSE)+D{r}*Inputs!$B$4")
        q.cell(r, 7, f"=E{r}*LaborRate")
        q.cell(r, 8, f"=F{r}+G{r}")
        q.cell(r, 9, f"=H{r}*1.15+Inputs!$B$6")
        q.cell(r, 10, f"=(I{r}-H{r})/I{r}")
        q.cell(r, 10).number_format = "0.0%"
    q["I13"] = "=H13*1.12+Inputs!$B$6"  # planted: inconsistent overhead
    q["G17"] = 3250                      # planted: hard-coded labor cost
    t = last + 1
    q.cell(t, 1, "Total")
    for col in "FGHI":
        q[f"{col}{t}"] = f"=SUM({col}{first}:{col}{last})"
    q[f"J{t}"] = f"=(I{t}-H{t})/I{t}"

    rep = wb.create_sheet("Report")
    rep["A1"], rep["B1"] = "Report date", "=TODAY()"
    rep["A3"], rep["B3"] = "Lowest margin", f"=MIN(Quote!J{first}:J{last})"
    rep["A4"], rep["B4"] = "Products below target", f'=COUNTIF(Quote!J{first}:J{last},"<"&Inputs!B7)'
    rep["A5"], rep["B5"] = "Selected cell", "J5"
    rep["A6"], rep["B6"] = "Selected margin", '=INDIRECT("Quote!"&B5)'

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    _freeze(path)
    return path


def _freeze(path: Path) -> None:
    """Pin the save timestamps openpyxl writes (zip entries, docProps), so the same sample always has
    the same bytes, and therefore the same workbook id in the web UI."""
    with zipfile.ZipFile(path) as z:
        entries = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            if name == "docProps/core.xml":
                data = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", b"2026-01-01T00:00:00Z", data)
            z.writestr(zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0)), data, zipfile.ZIP_DEFLATED)
