"""Score SheetIR table detection against the StatSheets benchmark (.xlsx sheets only).

StatSheets: 737 annotated sheets from 14 statistical agencies, with table ranges labeled by hand
(https://doi.org/10.5281/zenodo.21967189). Download and unzip it anywhere, then:

    uv run --with numpy python bench/eval_statsheets.py <path-to>/dataset [--jobs 4] [--out results.jsonl]

A predicted table spans its header (including group-header rows) through its total row. Predictions
are matched one-to-one to labeled tables by IoU (greedy, best first). Reported per IoU threshold:
precision, recall and F1 over all tables (micro), plus the share of sheets where every table matches.
"""

from __future__ import annotations

import argparse
import csv
import json
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

THRESHOLDS = (0.5, 0.8, 0.9, 1.0)


def detect(xlsx: str, sheet: str, scorer_spec: str = "default") -> list[str]:
    """Table ranges SheetIR detects on one sheet, as A1 ranges.

    `scorer_spec`: "default" (what build_ir uses), "hand", "model" (the packaged trained model), or a
    directory holding header_row.json and header_block.json."""
    from openpyxl.utils import get_column_letter

    from sheetir.build import WorkbookResolver
    from sheetir.formula import parse_formula
    from sheetir.groups import group_formulas
    from sheetir.loader import load
    from sheetir.structure import detect_tables
    from sheetir.tablemodel import ModelScorer, get_scorer

    scorer = (None if scorer_spec == "default" else get_scorer(scorer_spec) if scorer_spec in ("hand", "model")
              else ModelScorer.load(scorer_spec))

    raw = load(xlsx)
    resolver = WorkbookResolver(raw)
    parsed = {ck: parse_formula(c.formula, c.sheet, c.row, c.col, resolver)
              for ck, c in raw.cells.items() if c.formula is not None and c.sheet == sheet}
    groups = group_formulas({ck: pf.key for ck, pf in parsed.items()})
    tables = detect_tables(raw, {ck: pf.refs for ck, pf in parsed.items()},
                           {g.anchor for g in groups if g.axis == "col"}, scorer)
    return [f"{get_column_letter(t.c1)}{t.top}:{get_column_letter(t.c2)}{t.bottom}"
            for t in tables if t.sheet == sheet]


def box(a1: str) -> tuple[int, int, int, int]:
    from openpyxl.utils import range_boundaries
    c1, r1, c2, r2 = range_boundaries(a1)
    return r1, c1, r2, c2


def iou(a: str, b: str) -> float:
    (ar1, ac1, ar2, ac2), (br1, bc1, br2, bc2) = box(a), box(b)
    h = min(ar2, br2) - max(ar1, br1) + 1
    w = min(ac2, bc2) - max(ac1, bc1) + 1
    if h <= 0 or w <= 0:
        return 0.0
    inter = h * w
    return inter / ((ar2 - ar1 + 1) * (ac2 - ac1 + 1) + (br2 - br1 + 1) * (bc2 - bc1 + 1) - inter)


def match(pred: list[str], gold: list[str]) -> list[float]:
    """IoU of each one-to-one matched (pred, gold) pair, greedy by best IoU."""
    pairs = sorted(((iou(p, g), i, j) for i, p in enumerate(pred) for j, g in enumerate(gold)), reverse=True)
    used_p, used_g, out = set(), set(), []
    for v, i, j in pairs:
        if v > 0 and i not in used_p and j not in used_g:
            used_p.add(i)
            used_g.add(j)
            out.append(v)
    return out


def run_one(dataset: Path, row: dict, timeout: int, scorer_spec: str = "default") -> dict:
    import numpy as np
    gold = [g.decode() if isinstance(g, bytes) else str(g)
            for g in np.load(dataset / row["labels_path"], allow_pickle=False)["ranges_original"]]
    rec = {"file": row["file_path"], "sheet": row["sheet_name"], "source": row["source"], "gold": gold}
    t0 = time.perf_counter()
    try:
        proc = subprocess.run([sys.executable, __file__, "--one", str(dataset / row["file_path"]), row["sheet_name"], scorer_spec],
                              capture_output=True, text=True, timeout=timeout)
        if proc.returncode != 0:
            rec["error"] = (proc.stderr.strip().splitlines() or ["?"])[-1][:200]
            rec["pred"] = []
        else:
            rec["pred"] = json.loads(proc.stdout)
    except subprocess.TimeoutExpired:
        rec["error"], rec["pred"] = f"timeout {timeout}s", []
    rec["seconds"] = round(time.perf_counter() - t0, 2)
    rec["ious"] = match(rec["pred"], gold)
    return rec


def report(recs: list[dict]) -> None:
    n_gold = sum(len(r["gold"]) for r in recs)
    n_pred = sum(len(r["pred"]) for r in recs)
    errors = [r for r in recs if "error" in r]
    print(f"sheets {len(recs)}  labeled tables {n_gold}  predicted tables {n_pred}  "
          f"errors/timeouts {len(errors)}  median time {sorted(r['seconds'] for r in recs)[len(recs) // 2]}s")
    print(f"{'IoU >=':>8} {'precision':>10} {'recall':>8} {'F1':>6} {'sheets all-correct':>20}")
    for th in THRESHOLDS:
        tp = sum(sum(v >= th - 1e-9 for v in r["ious"]) for r in recs)
        p = tp / n_pred if n_pred else 0.0
        rc = tp / n_gold if n_gold else 0.0
        f1 = 2 * p * rc / (p + rc) if p + rc else 0.0
        exact = sum(len(r["pred"]) == len(r["gold"]) == sum(v >= th - 1e-9 for v in r["ious"]) for r in recs)
        print(f"{th:>8} {p:>10.1%} {rc:>8.1%} {f1:>6.1%} {exact / len(recs):>20.1%}")
    none = sum(1 for r in recs if not r["pred"] and "error" not in r)
    extra = sum(max(0, len(r["pred"]) - len(r["gold"])) for r in recs)
    print(f"sheets with no table found: {none}   surplus predictions: {extra}")
    by_src: dict[str, list] = {}
    for r in recs:
        by_src.setdefault(r["source"], []).append(r)
    print("\nby source (IoU >= 0.8 recall):")
    for src, rs in sorted(by_src.items()):
        g = sum(len(r["gold"]) for r in rs)
        tp = sum(sum(v >= 0.8 for v in r["ious"]) for r in rs)
        print(f"  {src:<24} {len(rs):>4} sheets  {tp / g if g else 0:>6.1%}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("dataset", type=Path, help="StatSheets dataset/ directory (holds manifest.csv)")
    ap.add_argument("--jobs", type=int, default=4)
    ap.add_argument("--timeout", type=int, default=180)
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--out", type=Path)
    ap.add_argument("--scorer", default="default", help='"default", "hand", "model", or a model directory')
    args = ap.parse_args()

    with open(args.dataset / "manifest.csv", encoding="utf-8") as f:
        rows = [r for r in csv.DictReader(f) if r["file_path"].lower().endswith(".xlsx")]
    if args.limit:
        rows = rows[:args.limit]
    with ThreadPoolExecutor(args.jobs) as pool:
        recs = list(pool.map(lambda r: run_one(args.dataset, r, args.timeout, args.scorer), rows))
    if args.out:
        args.out.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in recs) + "\n", encoding="utf-8")
    report(recs)


if __name__ == "__main__":
    if len(sys.argv) in (4, 5) and sys.argv[1] == "--one":
        print(json.dumps(detect(*sys.argv[2:])))
    else:
        main()
